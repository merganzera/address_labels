#!/usr/bin/env python3
"""Generate Avery 18660 label PDFs from an address spreadsheet."""

import argparse
import os
import re
from typing import Dict, Iterable, List, Optional


LABEL_SPEC = {
    "page_width_in": 8.5,
    "page_height_in": 11.0,
    "columns": 3,
    "rows": 10,
    "label_width_in": 2.625,
    "label_height_in": 1.0,
    "left_margin_in": 0.1875,
    "top_margin_in": 0.5,
    "horizontal_pitch_in": 2.75,
    "vertical_pitch_in": 1.0,
    "text_padding_in": 0.1,
}

COLUMN_ALIASES = {
    "addressee": ["addressee", "recipient", "name", "full name"],
    "first_name": ["first name", "firstname", "first"],
    "last_name": ["last name", "lastname", "last"],
    "address1": ["address 1", "address1", "street", "street address", "street1"],
    "address2": ["address 2", "address2", "unit", "apt", "suite", "street2"],
    "city": ["city", "town"],
    "state": ["state", "province", "region"],
    "postal": ["zip", "zip code", "postal", "postal code"],
    "country": ["country", "country code"],
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create Avery 18660 label PDFs from an XLSX address list."
    )
    parser.add_argument(
        "--input",
        required=True,
        help="Path to the XLSX file containing addresses.",
    )
    parser.add_argument(
        "--output",
        required=True,
        help="Path to the output PDF file.",
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Optional worksheet name. Defaults to the first sheet.",
    )
    parser.add_argument(
        "--font",
        default=None,
        help="Path to a specific font file (TTF/OTF). Overrides --font-name.",
    )
    parser.add_argument(
        "--font-name",
        action="append",
        default=[],
        help="Preferred font name (repeat or comma-separate). Will search OSX Font Book.",
    )
    parser.add_argument(
        "--font-size",
        type=float,
        default=12.0,
        help="Font size for label text.",
    )
    parser.add_argument(
        "--leading",
        type=float,
        default=14.0,
        help="Line spacing in points.",
    )
    parser.add_argument(
        "--fit-mode",
        choices=["wrap", "shrink"],
        default="shrink",
        help="When text is too wide, either wrap lines or shrink the label font.",
    )
    return parser.parse_args()


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    return "".join(ch.lower() for ch in str(value) if ch.isalnum())


def find_header_index(header_map: Dict[str, int], aliases: Iterable[str]) -> Optional[int]:
    for alias in aliases:
        idx = header_map.get(normalize_header(alias))
        if idx is not None:
            return idx
    return None


def clean_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        return str(value)
    text = str(value)
    text = re.sub(r"[\x00-\x1F\x7F]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def format_postal(value: object) -> str:
    text = clean_cell(value)
    if text.isdigit() and len(text) < 5:
        return text.zfill(5)
    return text


def format_city_state_zip(city: str, state: str, postal: str) -> str:
    parts = []
    if city:
        parts.append(city)
    if state:
        if parts:
            parts[-1] = f"{parts[-1]}, {state}"
        else:
            parts.append(state)
    if postal:
        if parts:
            parts[-1] = f"{parts[-1]} {postal}"
        else:
            parts.append(postal)
    return parts[0] if parts else ""


def is_us_country(value: str) -> bool:
    normalized = "".join(ch.lower() for ch in value.strip() if ch.isalnum())
    return normalized in {"us", "usa", "unitedstates", "unitedstatesofamerica"}


def build_lines(entry: Dict[str, str]) -> List[tuple[str, str]]:
    lines: List[tuple[str, str]] = []
    if entry.get("addressee"):
        lines.append(("addressee", entry["addressee"]))
    if entry.get("address1"):
        lines.append(("address1", entry["address1"]))
    if entry.get("address2"):
        lines.append(("address2", entry["address2"]))
    city_state_zip = format_city_state_zip(
        entry.get("city", ""),
        entry.get("state", ""),
        entry.get("postal", ""),
    )
    if city_state_zip:
        lines.append(("city_state_zip", city_state_zip))
    country = entry.get("country", "")
    if country and not is_us_country(country):
        lines.append(("country", country.upper()))
    return lines


def load_addresses(path: str, sheet_name: Optional[str]) -> List[List[tuple[str, str]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit(
            "Missing dependency: openpyxl. Install with 'pip install openpyxl'."
        ) from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active

    header_cells = next(worksheet.iter_rows(min_row=1, max_row=1))
    headers = [normalize_header(cell.value) for cell in header_cells]
    header_map = {header: idx for idx, header in enumerate(headers) if header}

    indices = {
        key: find_header_index(header_map, aliases)
        for key, aliases in COLUMN_ALIASES.items()
    }

    entries: List[List[tuple[str, str]]] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        values = list(row)
        addressee = clean_cell(values[indices["addressee"]]) if indices["addressee"] is not None else ""
        first_name = clean_cell(values[indices["first_name"]]) if indices["first_name"] is not None else ""
        last_name = clean_cell(values[indices["last_name"]]) if indices["last_name"] is not None else ""
        if not addressee:
            addressee = " ".join(part for part in [first_name, last_name] if part).strip()

        entry = {
            "addressee": addressee,
            "address1": clean_cell(values[indices["address1"]]) if indices["address1"] is not None else "",
            "address2": clean_cell(values[indices["address2"]]) if indices["address2"] is not None else "",
            "city": clean_cell(values[indices["city"]]) if indices["city"] is not None else "",
            "state": clean_cell(values[indices["state"]]) if indices["state"] is not None else "",
            "postal": format_postal(values[indices["postal"]]) if indices["postal"] is not None else "",
            "country": clean_cell(values[indices["country"]]) if indices["country"] is not None else "",
        }

        if not any(entry.values()):
            continue

        lines = build_lines(entry)
        if lines:
            entries.append(lines)

    return entries


def normalize_font_name(value: str) -> str:
    return "".join(ch.lower() for ch in value if ch.isalnum())


def iter_font_files() -> Iterable[str]:
    roots = [
        os.path.expanduser("~/Library/Fonts"),
        "/Library/Fonts",
        "/System/Library/Fonts",
        "/System/Library/Fonts/Supplemental",
        "/usr/share/fonts",
        "/usr/local/share/fonts",
    ]
    for root in roots:
        if not os.path.isdir(root):
            continue
        for dirpath, _, filenames in os.walk(root):
            for filename in filenames:
                lower_name = filename.lower()
                if lower_name.endswith((".ttf", ".otf")):
                    yield os.path.join(dirpath, filename)


def find_font_path(font_name: str) -> Optional[str]:
    if not font_name:
        return None
    if os.path.isfile(font_name):
        return font_name
    normalized = normalize_font_name(font_name)
    partial_matches = []
    for path in iter_font_files():
        base_name = os.path.splitext(os.path.basename(path))[0]
        base_norm = normalize_font_name(base_name)
        if base_norm == normalized:
            return path
        if normalized in base_norm:
            partial_matches.append(path)
    return partial_matches[0] if partial_matches else None


def expand_font_names(raw_names: Iterable[str]) -> List[str]:
    names: List[str] = []
    for item in raw_names:
        if not item:
            continue
        for part in item.split(","):
            name = part.strip()
            if name:
                names.append(name)
    return names


def resolve_font(font_path: Optional[str], font_names: Iterable[str]) -> str:
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError as exc:
        raise SystemExit(
            "Missing dependency: reportlab. Install with 'pip install reportlab'."
        ) from exc

    def register_font(path: str, label: str) -> str:
        pdfmetrics.registerFont(TTFont(label, path))
        return label

    if font_path:
        if not os.path.isfile(font_path):
            raise SystemExit(f"Font file not found: {font_path}")
        label = os.path.splitext(os.path.basename(font_path))[0] or "CustomFont"
        return register_font(font_path, label)

    candidates = expand_font_names(font_names)
    candidates.append("Jost")

    for name in candidates:
        if name in pdfmetrics.standardFonts:
            return name
        path = find_font_path(name)
        if path:
            label = name.replace(" ", "") or os.path.splitext(os.path.basename(path))[0]
            return register_font(path, label)

    return "Helvetica"


def find_split_candidates(text: str, kind: str) -> List[tuple[int, int]]:
    candidates: List[tuple[int, int]] = []

    def add_matches(pattern: str, priority: int, use_end: bool = False) -> None:
        for match in re.finditer(pattern, text, flags=re.IGNORECASE):
            idx = match.end() if use_end else match.start()
            candidates.append((priority, idx))

    if kind == "addressee":
        add_matches(r",\s+", 1, use_end=True)
        add_matches(r"\s+and\s+", 1)
        add_matches(r"\s+&\s+", 1)
    elif kind in {"address1", "address2"}:
        for token in [
            "apt",
            "apartment",
            "unit",
            "suite",
            "ste",
            "floor",
            "fl",
            "bldg",
            "building",
            "dept",
            "attn",
            "c/o",
            "#",
        ]:
            add_matches(rf"\s+{re.escape(token)}\b", 1)
        add_matches(r",\s+", 2, use_end=True)
    elif kind == "city_state_zip":
        add_matches(r",\s+", 1, use_end=True)
        add_matches(r"\s+\d", 2)

    add_matches(r"\s+", 3)
    return candidates


def split_line(
    text: str,
    kind: str,
    max_width: float,
    font_name: str,
    font_size: float,
) -> List[str]:
    from reportlab.pdfbase import pdfmetrics

    if not text:
        return []

    if pdfmetrics.stringWidth(text, font_name, font_size) <= max_width:
        return [text]

    candidates = find_split_candidates(text, kind)
    best = None
    best_score = None

    for priority, idx in candidates:
        left = text[:idx].strip()
        right = text[idx:].strip()
        if not left or not right:
            continue
        left_width = pdfmetrics.stringWidth(left, font_name, font_size)
        right_width = pdfmetrics.stringWidth(right, font_name, font_size)
        fits = left_width <= max_width and right_width <= max_width
        score = (0 if fits else 1, priority, max(left_width, right_width), abs(left_width - right_width))
        if best_score is None or score < best_score:
            best_score = score
            best = (left, right)

    if best:
        return [best[0], best[1]]

    # Last resort: split the word by width.
    for idx in range(len(text) - 1, 0, -1):
        left = text[:idx].strip()
        right = text[idx:].strip()
        if not left or not right:
            continue
        if pdfmetrics.stringWidth(left, font_name, font_size) <= max_width:
            return [left, right]

    return [text]


def wrap_lines(
    lines: List[tuple[str, str]],
    max_width: float,
    font_name: str,
    font_size: float,
) -> List[str]:
    wrapped: List[str] = []
    for kind, text in lines:
        wrapped.extend(split_line(text, kind, max_width, font_name, font_size))
    return wrapped


def draw_labels(
    output_path: str,
    labels: List[List[tuple[str, str]]],
    font_name: str,
    font_size: float,
    leading: float,
    fit_mode: str,
) -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import inch

    page_width = LABEL_SPEC["page_width_in"] * inch
    page_height = LABEL_SPEC["page_height_in"] * inch
    label_width = LABEL_SPEC["label_width_in"] * inch
    label_height = LABEL_SPEC["label_height_in"] * inch
    left_margin = LABEL_SPEC["left_margin_in"] * inch
    top_margin = LABEL_SPEC["top_margin_in"] * inch
    horizontal_pitch = LABEL_SPEC["horizontal_pitch_in"] * inch
    vertical_pitch = LABEL_SPEC["vertical_pitch_in"] * inch
    text_padding = LABEL_SPEC["text_padding_in"] * inch
    columns = LABEL_SPEC["columns"]
    rows = LABEL_SPEC["rows"]
    labels_per_page = columns * rows
    max_text_width = label_width - 2 * text_padding

    pdf = canvas.Canvas(output_path, pagesize=(page_width, page_height))

    for index, lines in enumerate(labels):
        if index and index % labels_per_page == 0:
            pdf.showPage()

        slot = index % labels_per_page
        row = slot // columns
        col = slot % columns

        origin_x = left_margin + col * horizontal_pitch
        origin_y = page_height - top_margin - label_height - row * vertical_pitch

        if not lines:
            continue

        if fit_mode == "wrap":
            rendered_lines = wrap_lines(lines, max_text_width, font_name, font_size)
            label_font_size = font_size
            label_leading = leading
        else:
            rendered_lines = [text for _, text in lines if text]
            max_line_width = max(
                (pdfmetrics.stringWidth(text, font_name, font_size) for text in rendered_lines),
                default=0.0,
            )
            scale = 1.0
            if max_line_width > max_text_width and max_line_width > 0:
                scale = max_text_width / max_line_width
            label_font_size = font_size * scale
            label_leading = leading * scale

        if not rendered_lines:
            continue

        pdf.setFont(font_name, label_font_size)
        block_height = (len(rendered_lines) - 1) * label_leading + label_font_size
        baseline_y = origin_y + (label_height + block_height) / 2 - label_font_size

        for line in rendered_lines:
            text_width = pdfmetrics.stringWidth(line, font_name, label_font_size)
            text_x = origin_x + (label_width - text_width) / 2
            pdf.drawString(text_x, baseline_y, line)
            baseline_y -= label_leading

    pdf.save()


def main() -> None:
    args = parse_args()

    if not os.path.isfile(args.input):
        raise SystemExit(f"Input file not found: {args.input}")

    labels = load_addresses(args.input, args.sheet)
    if not labels:
        raise SystemExit("No address rows found to render.")

    font_name = resolve_font(args.font, args.font_name)
    draw_labels(args.output, labels, font_name, args.font_size, args.leading, args.fit_mode)

    print(f"Wrote {len(labels)} labels to {args.output}.")


if __name__ == "__main__":
    main()
